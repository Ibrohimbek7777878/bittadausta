"""client_erp/views/admin.py — Mini ERP Admin panel (MebelCity xodimlari uchun)."""
import random
import re
import json
import string
from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse
from django.views.decorators.http import require_POST
from ..models import (
    ClientUser, ClientVIPLevel,
    GamificationRule, Announcement, MonthlyDiscount,
    Event, EventParticipant, RewardItem, RewardClaim,
    Quest, ClientOrderStatusDef,
    ClientPlan, ClientSubscription,
    CoinPack, ClientAIPrice,
)
from .auth import _normalize_phone
from ..models import ClientLoginAttempt


def _staff_required(view_func):
    """Staff-only decorator."""
    from functools import wraps
    @wraps(view_func)
    @login_required(login_url='login_e')
    def wrapper(request, *args, **kwargs):
        if not request.user.is_staff:
            return JsonResponse({'error': 'Ruxsat yo\'q'}, status=403)
        return view_func(request, *args, **kwargs)
    return wrapper


@_staff_required
def bom_settings_page(request):
    """BOM Narx standarti — GLOBAL standart, faqat admin panel orqali (is_staff). /mini/bom-settings/."""
    return render(request, 'client_erp/bom_settings.html', {'user': request.user})


@_staff_required
def mini_admin_dashboard(request):
    """GET /mini/admin/ — admin dashboard."""
    users = ClientUser.objects.select_related('vip_level').order_by('-created_at')
    levels = ClientVIPLevel.objects.all().order_by('level_number')
    rules = GamificationRule.objects.all().order_by('sort_order')
    announcements = Announcement.objects.order_by('-created_at')[:20]
    events = Event.objects.order_by('-event_date')[:10]
    rewards = RewardItem.objects.order_by('sort_order')
    quests = Quest.objects.order_by('quest_type', 'sort_order')
    pending_claims = RewardClaim.objects.filter(status='pending').select_related('user', 'reward')
    order_statuses = ClientOrderStatusDef.objects.all().order_by('sort_order', 'id')

    # Statistika
    from django.db.models import Sum, Count
    stats = {
        'total_users': users.count(),
        'active_users': users.filter(is_active=True, is_blocked=False).count(),
        'verified_users': users.filter(is_verified=True).count(),
        'total_xp': users.aggregate(t=Sum('xp'))['t'] or 0,
        'total_coins': users.aggregate(t=Sum('coins'))['t'] or 0,
        'pending_claims': pending_claims.count(),
    }

    # Oxirgi 7 kunda ro'yxatdan o'tgan userlar (eng so'nggi birinchi).
    # 2026-07-16: ilgari faqat BUGUN (ko'pincha 0) ko'rsatilardi — foydasiz edi.
    # Endi «Oxirgi hafta (7 kun)» — users allaqachon -created_at tartibda.
    from django.utils import timezone as _tz
    from datetime import timedelta as _td
    _today = _tz.localdate()
    _week_start = _today - _td(days=6)          # bugun + oldingi 6 kun = 7 kun
    week_users = users.filter(created_at__date__gte=_week_start)

    # Impersonate holati
    imp_user_id = request.session.get('impersonating_client_id')
    imp_user = None
    if imp_user_id:
        imp_user = ClientUser.objects.filter(pk=imp_user_id).first()

    return render(request, 'client_erp/admin/dashboard.html', {
        'client_users': users,
        'week_users': week_users,
        'week_count': week_users.count(),
        'week_start': _week_start,
        'levels': levels,
        'rules': rules,
        'announcements': announcements,
        'events': events,
        'rewards': rewards,
        'quests': quests,
        'pending_claims': pending_claims,
        'stats': stats,
        'imp_user': imp_user,
        'order_statuses': order_statuses,
    })


@login_required(login_url='login_e')
@require_POST
def mini_admin_create_user(request):
    """POST /mini/admin/create-user/"""
    if not request.user.is_staff:
        return JsonResponse({'error': 'Ruxsat yo\'q'}, status=403)

    body = json.loads(request.body)
    phone = _normalize_phone((body.get('phone', '') or '').strip())
    name = (body.get('full_name', '') or '').strip()
    org = (body.get('organization', '') or '').strip()

    if not phone or not name:
        return JsonResponse({'error': 'Telefon va ism kerak'}, status=400)

    if ClientUser.objects.filter(phone=phone).exists():
        return JsonResponse({'error': 'Bu telefon allaqachon ro\'yxatdan o\'tgan'}, status=400)

    password = str(random.randint(100000, 999999))
    username = ClientUser.generate_username(name)

    user = ClientUser(phone=phone, username=username, full_name=name, organization=org)
    user.set_password(password)
    user.is_verified = True
    user.save()

    # Katta ERP Client bilan bog'lash (telefon bo'yicha)
    try:
        from clients.models import Client
        from django.db.models import Q
        import re
        digits = re.sub(r'\D', '', phone)[-9:]
        client = Client.objects.filter(Q(phone__endswith=digits) | Q(phone2__endswith=digits)).first()
        if client:
            user.client = client
            user.save(update_fields=['client'])
    except Exception:
        pass

    return JsonResponse({
        'ok': True,
        'id': user.id,
        'username': user.username,
        'password': password,
        'message': f'Yaratildi! Login: {phone}, Parol: {password}',
    })


@login_required(login_url='login_e')
@require_POST
def mini_admin_generate_password(request, user_id):
    """POST /mini/admin/password/<id>/ — parol o'rnatish.
    Body'da 'password' bo'lsa — o'sha qo'lda parol saqlanadi (admin tahrirlagan);
    bo'lmasa — tasodifiy 6-xonali generatsiya qilinadi (backward-compat)."""
    if not request.user.is_staff:
        return JsonResponse({'error': 'Ruxsat yo\'q'}, status=403)

    user = get_object_or_404(ClientUser, pk=user_id)
    try:
        body = json.loads(request.body or '{}')
    except (ValueError, TypeError):
        body = {}
    custom = (body.get('password') or '').strip()
    if custom:
        if len(custom) < 4:
            return JsonResponse({'error': 'Parol kamida 4 ta belgi bo\'lishi kerak'}, status=400)
        password = custom
    else:
        password = str(random.randint(100000, 999999))
    user.set_password(password)
    user.save(update_fields=['password_hash'])

    return JsonResponse({'ok': True, 'password': password, 'message': f'{user.full_name} uchun parol saqlandi: {password}'})


@login_required(login_url='login_e')
@require_POST
def mini_admin_toggle(request, user_id):
    """POST /mini/admin/toggle/<id>/ — verify/block/unblock."""
    if not request.user.is_staff:
        return JsonResponse({'error': 'Ruxsat yo\'q'}, status=403)

    body = json.loads(request.body)
    action = body.get('action')
    user = get_object_or_404(ClientUser, pk=user_id)

    if action == 'verify':
        user.is_verified = True
        user.save(update_fields=['is_verified'])
    elif action == 'block':
        user.is_blocked = True
        user.save(update_fields=['is_blocked'])
    elif action == 'unblock':
        user.is_blocked = False
        user.save(update_fields=['is_blocked'])
    elif action == 'unverify':
        user.is_verified = False
        user.save(update_fields=['is_verified'])
    else:
        return JsonResponse({'error': 'Noto\'g\'ri action'}, status=400)

    return JsonResponse({'ok': True})


@login_required(login_url='login_e')
@require_POST
def mini_admin_delete_user(request, user_id):
    """POST /mini/admin/delete-user/<id>/ — akkountni BUTUNLAY o'chirish (admin).

    Ichki client_erp ma'lumoti + disk fayllari o'chadi; katta ERP va panoramalar
    teginilmaydi (purge_client_account — ORM SET_NULL kafolati)."""
    if not request.user.is_staff:
        return JsonResponse({'error': 'Ruxsat yo\'q'}, status=403)

    # Foydalanuvchi mavjudligini joriy tenant DB'da tekshirish
    user = get_object_or_404(ClientUser, pk=user_id)
    from ..services.account_purge import purge_client_account
    from tenant_manager.middleware import get_current_db_alias
    db = get_current_db_alias()
    res = purge_client_account(user.pk, db=db)
    if not res.get('ok'):
        return JsonResponse({'error': res.get('error') or 'O\'chirishda xatolik'}, status=500)
    return JsonResponse({'ok': True, 'files_deleted': res.get('files_deleted', 0)})


# ═══════════════════════════════════════════════════════════════════
#  IMPERSONATE — mijoz hisobiga kirish
# ═══════════════════════════════════════════════════════════════════

@_staff_required
@require_POST
def mini_admin_impersonate(request, user_id):
    """POST /mini/admin/impersonate/<id>/ — mijoz hisobiga kirish.

    ClientERPMiddleware Django session'ga tegmaydi, faqat JWT cookie'ni
    tekshiradi ([[client_erp/middleware.py]]) — shu sabab bu yerda ham
    haqiqiy client_erp_token cookie o'rnatilishi SHART, aks holda "Kirish"
    tugmasi faqat foydasiz session-flag qo'yib, brauzerning eski cookie'siga
    (masalan boshqa mijoz hisobiga) qaytarib yuboradi.
    """
    from ..auth_backend import generate_token, TOKEN_COOKIE
    user = get_object_or_404(ClientUser, pk=user_id)
    request.session['impersonating_client_id'] = user.id
    token = generate_token(user)
    resp = JsonResponse({'ok': True, 'redirect': f'/mini/{user.username}/'})
    resp.set_cookie(TOKEN_COOKIE, token, max_age=30 * 24 * 3600, httponly=True, samesite='Lax')
    return resp


@_staff_required
def mini_admin_stop_impersonate(request):
    """GET /mini/admin/stop-impersonate/ — chiqish."""
    request.session.pop('impersonating_client_id', None)
    return redirect('client_erp:admin')


# ═══════════════════════════════════════════════════════════════════
#  GAMIFICATION RULES CRUD
# ═══════════════════════════════════════════════════════════════════

@_staff_required
@require_POST
def mini_admin_save_rule(request):
    """POST /mini/admin/save-rule/ — qoida yaratish/yangilash."""
    body = json.loads(request.body)
    rule_id = body.get('id')
    if rule_id:
        rule = get_object_or_404(GamificationRule, pk=rule_id)
    else:
        rule = GamificationRule()

    rule.code = body.get('code', rule.code or '')
    rule.name = body.get('name', '')
    rule.description = body.get('description', '')
    rule.icon = body.get('icon', '⭐')
    rule.xp_reward = int(body.get('xp_reward', 0))
    rule.coin_reward = int(body.get('coin_reward', 0))
    rule.is_active = bool(body.get('is_active', True))
    rule.is_repeatable = bool(body.get('is_repeatable', True))
    rule.max_per_day = body.get('max_per_day') or None
    rule.sort_order = int(body.get('sort_order', 0))
    rule.save()
    return JsonResponse({'ok': True, 'id': rule.id})


@_staff_required
@require_POST
def mini_admin_delete_rule(request, pk):
    """POST /mini/admin/delete-rule/<pk>/"""
    GamificationRule.objects.filter(pk=pk).delete()
    return JsonResponse({'ok': True})


# ═══════════════════════════════════════════════════════════════════
#  BUYURTMA STATUSLARI CRUD (ClientOrderStatusDef)
# ═══════════════════════════════════════════════════════════════════

@_staff_required
@require_POST
def mini_admin_save_status(request):
    """POST /mini/admin/save-status/ — status yaratish/yangilash.

    Tizim statuslarining (is_system=True) `key`i o'zgartirilmaydi — faqat
    label/color/sort_order/is_active. Yangi status uchun key SlugField orqali
    avtomatik yasaladi (label'dan) va noyob bo'lishi tekshiriladi.
    """
    from django.utils.text import slugify

    body = json.loads(request.body)
    status_id = body.get('id')
    label = (body.get('label') or '').strip()
    if not label:
        return JsonResponse({'ok': False, 'error': 'Nomi majburiy'}, status=400)

    if status_id:
        status = get_object_or_404(ClientOrderStatusDef, pk=status_id)
    else:
        key = slugify(label, allow_unicode=False).replace('-', '_')[:30] or 'status'
        base_key, i = key, 2
        while ClientOrderStatusDef.objects.filter(key=key).exists():
            key = f'{base_key}_{i}'[:30]
            i += 1
        status = ClientOrderStatusDef(key=key)

    status.label = label
    status.color = body.get('color') or status.color or '#6366f1'
    status.sort_order = int(body.get('sort_order', status.sort_order or 0))
    status.is_active = bool(body.get('is_active', True))
    status.save()
    return JsonResponse({'ok': True, 'id': status.id, 'key': status.key})


@_staff_required
@require_POST
def mini_admin_delete_status(request, pk):
    """POST /mini/admin/delete-status/<pk>/ — tizim statuslari o'chirilmaydi."""
    status = get_object_or_404(ClientOrderStatusDef, pk=pk)
    if status.is_system:
        return JsonResponse({'ok': False, 'error': "Tizim statusini o'chirib bo'lmaydi — faqat yashirish (Faol=yo'q)"}, status=400)
    status.delete()
    return JsonResponse({'ok': True})


# ═══════════════════════════════════════════════════════════════════
#  ANNOUNCEMENTS CRUD
# ═══════════════════════════════════════════════════════════════════

@_staff_required
@require_POST
def mini_admin_save_announcement(request):
    """POST /mini/admin/save-announcement/"""
    body = json.loads(request.body)
    ann_id = body.get('id')
    if ann_id:
        ann = get_object_or_404(Announcement, pk=ann_id)
    else:
        ann = Announcement(created_by=request.user)

    from django.utils.dateparse import parse_datetime
    ann.announcement_type = body.get('type', 'news')
    ann.title = body.get('title', '')
    ann.body = body.get('body', '')
    ann.target_all = bool(body.get('target_all', True))
    ann.discount_percent = body.get('discount_percent') or None
    ann.discount_amount = body.get('discount_amount') or None
    ann.discount_code = body.get('discount_code', '')
    ann.start_date = parse_datetime(body.get('start_date', '')) or __import__('django.utils.timezone', fromlist=['now']).now()
    end = body.get('end_date')
    ann.end_date = parse_datetime(end) if end else None
    ann.is_active = bool(body.get('is_active', True))
    ann.is_pinned = bool(body.get('is_pinned', False))
    ann.send_telegram = bool(body.get('send_telegram', False))
    ann.save()

    # Target levels
    levels = body.get('target_level_ids', [])
    if levels:
        ann.target_levels.set(levels)
        ann.target_all = False
        ann.save(update_fields=['target_all'])

    return JsonResponse({'ok': True, 'id': ann.id})


@_staff_required
@require_POST
def mini_admin_delete_announcement(request, pk):
    Announcement.objects.filter(pk=pk).delete()
    return JsonResponse({'ok': True})


# ═══════════════════════════════════════════════════════════════════
#  VIP LEVELS
# ═══════════════════════════════════════════════════════════════════

@_staff_required
@require_POST
def mini_admin_save_level(request):
    """POST /mini/admin/save-level/"""
    body = json.loads(request.body)
    lvl_id = body.get('id')
    if lvl_id:
        lvl = get_object_or_404(ClientVIPLevel, pk=lvl_id)
    else:
        lvl = ClientVIPLevel()

    lvl.level_number = int(body.get('level_number', 1))
    lvl.name = body.get('name', '')
    lvl.min_turnover = body.get('min_turnover', 0)
    lvl.discount_percent = body.get('discount_percent', 0)
    lvl.referral_commission = body.get('referral_commission', 0)
    lvl.color = body.get('color', '#cd7f32')
    lvl.badge_icon = body.get('badge_icon', '⭐')
    lvl.description = body.get('description', '')
    lvl.auto_promote = bool(body.get('auto_promote', True))
    lvl.require_admin_approval = bool(body.get('require_admin_approval', False))
    lvl.downgrade_after_months = int(body.get('downgrade_after_months', 6))
    lvl.save()
    return JsonResponse({'ok': True, 'id': lvl.id})


# ═══════════════════════════════════════════════════════════════════
#  EVENTS CRUD
# ═══════════════════════════════════════════════════════════════════

@_staff_required
@require_POST
def mini_admin_save_event(request):
    """POST /mini/admin/save-event/"""
    body = json.loads(request.body)
    from django.utils.dateparse import parse_datetime
    evt_id = body.get('id')
    if evt_id:
        evt = get_object_or_404(Event, pk=evt_id)
    else:
        evt = Event(created_by=request.user)

    evt.title = body.get('title', '')
    evt.description = body.get('description', '')
    evt.event_date = parse_datetime(body.get('event_date', '')) or __import__('django.utils.timezone', fromlist=['now']).now()
    evt.location = body.get('location', '')
    evt.location_url = body.get('location_url', '')
    evt.max_participants = int(body.get('max_participants', 0))
    evt.xp_reward = int(body.get('xp_reward', 500))
    evt.coin_reward = int(body.get('coin_reward', 150))
    evt.target_all = bool(body.get('target_all', True))
    evt.is_active = bool(body.get('is_active', True))
    evt.save()
    return JsonResponse({'ok': True, 'id': evt.id, 'qr_code': evt.qr_code})


@_staff_required
@require_POST
def mini_admin_delete_event(request, pk):
    Event.objects.filter(pk=pk).delete()
    return JsonResponse({'ok': True})


# ═══════════════════════════════════════════════════════════════════
#  REWARDS CRUD
# ═══════════════════════════════════════════════════════════════════

@_staff_required
@require_POST
def mini_admin_save_reward(request):
    """POST /mini/admin/save-reward/"""
    body = json.loads(request.body)
    rw_id = body.get('id')
    if rw_id:
        rw = get_object_or_404(RewardItem, pk=rw_id)
    else:
        rw = RewardItem()

    rw.name = body.get('name', '')
    rw.description = body.get('description', '')
    rw.reward_type = body.get('reward_type', 'coupon')
    rw.coin_price = int(body.get('coin_price', 0))
    rw.discount_value = body.get('discount_value') or None
    rw.stock = int(body.get('stock', 0))
    rw.is_active = bool(body.get('is_active', True))
    rw.sort_order = int(body.get('sort_order', 0))
    rw.save()
    return JsonResponse({'ok': True, 'id': rw.id})


@_staff_required
@require_POST
def mini_admin_claim_action(request, pk):
    """POST /mini/admin/claim/<pk>/ — tasdiqlash/bekor qilish."""
    body = json.loads(request.body)
    action = body.get('action')  # 'approve' | 'deliver' | 'cancel'
    claim = get_object_or_404(RewardClaim, pk=pk)

    from django.utils import timezone
    if action == 'approve':
        claim.status = 'approved'
        claim.approved_by = request.user
        claim.approved_at = timezone.now()
    elif action == 'deliver':
        claim.status = 'delivered'
    elif action == 'cancel':
        claim.status = 'cancelled'
        # Tangani qaytarish
        claim.user.coins += claim.coins_spent
        claim.user.save(update_fields=['coins'])
    claim.save()
    return JsonResponse({'ok': True})


# ═══════════════════════════════════════════════════════════════════
#  QUESTS CRUD
# ═══════════════════════════════════════════════════════════════════

@_staff_required
@require_POST
def mini_admin_save_quest(request):
    """POST /mini/admin/save-quest/"""
    body = json.loads(request.body)
    from django.utils.dateparse import parse_date
    q_id = body.get('id')
    if q_id:
        q = get_object_or_404(Quest, pk=q_id)
    else:
        q = Quest()

    q.quest_type = body.get('quest_type', 'daily')
    q.title = body.get('title', '')
    q.description = body.get('description', '')
    q.icon = body.get('icon', '📋')
    q.xp_reward = int(body.get('xp_reward', 0))
    q.coin_reward = int(body.get('coin_reward', 0))
    q.action_type = body.get('action_type', 'custom')
    q.action_count = int(body.get('action_count', 1))
    q.start_date = parse_date(body.get('start_date', '')) if body.get('start_date') else None
    q.end_date = parse_date(body.get('end_date', '')) if body.get('end_date') else None
    q.is_active = bool(body.get('is_active', True))
    q.sort_order = int(body.get('sort_order', 0))
    q.save()
    return JsonResponse({'ok': True, 'id': q.id})


@_staff_required
@require_POST
def mini_admin_delete_quest(request, pk):
    Quest.objects.filter(pk=pk).delete()
    return JsonResponse({'ok': True})


@_staff_required
@require_POST
def mini_admin_delete_reward(request, pk):
    RewardItem.objects.filter(pk=pk).delete()
    return JsonResponse({'ok': True})


# ═══════════════════════════════════════════════════════════════════
#  VISIT CARDS — helpers & bulk endpoints
# ═══════════════════════════════════════════════════════════════════

def _gen_password():
    """MC-XXXX formatda parol generatsiya qilish."""
    return 'MC-' + ''.join(random.choices(string.ascii_uppercase + string.digits, k=4))


def _get_card_status(user):
    """VIP level → status string."""
    if not user.vip_level:
        return 'start'
    n = user.vip_level.level_number
    if n == 1:
        return 'hamkor'
    elif n == 2:
        return 'silver'
    elif n == 3:
        return 'gold'
    else:
        return 'vip'


@_staff_required
@require_POST
def mini_admin_bulk_check(request):
    """POST /mini/admin/bulk-check/ — telefonlar ro'yxatini tekshirish."""
    body = json.loads(request.body)
    rows = body.get('rows', [])
    results = []
    summary = {'new': 0, 'exists': 0, 'error': 0}

    for row in rows:
        phone = _normalize_phone(str(row.get('phone', '')))
        org = (row.get('org', '') or '').strip()
        name = (row.get('name', '') or '').strip()

        digits = re.sub(r'\D', '', phone)
        if len(digits) != 12 or not digits.startswith('998'):
            results.append({
                'phone': row.get('phone', ''),
                'org': org,
                'name': name,
                'state': 'error',
                'error': 'Telefon raqam noto\'g\'ri (12 ta raqam, 998 bilan boshlanishi kerak)',
            })
            summary['error'] += 1
            continue

        # DB tekshirish
        existing = ClientUser.objects.filter(phone=phone).first()
        if existing:
            results.append({
                'phone': phone,
                'org': org,
                'name': name,
                'state': 'exists',
                'user_id': existing.id,
                'username': existing.username,
            })
            summary['exists'] += 1
        else:
            results.append({
                'phone': phone,
                'org': org,
                'name': name,
                'state': 'new',
            })
            summary['new'] += 1

    return JsonResponse({'ok': True, 'results': results, 'summary': summary})


@_staff_required
@require_POST
def mini_admin_bulk_create(request):
    """POST /mini/admin/bulk-create/ — ommaviy foydalanuvchi yaratish."""
    body = json.loads(request.body)
    rows = body.get('rows', [])
    status_str = body.get('status', 'start')
    include_existing = body.get('include_existing', True)

    # Status → VIP level mapping
    status_level_map = {
        'start': None,
        'hamkor': 1,
        'silver': 2,
        'gold': 3,
        'vip': 4,
    }
    target_level_number = status_level_map.get(status_str)
    target_level = None
    if target_level_number is not None:
        target_level = ClientVIPLevel.objects.filter(level_number=target_level_number).first()

    users_out = []
    created = 0
    updated = 0

    for row in rows:
        phone = _normalize_phone(str(row.get('phone', '')))
        org = (row.get('org', '') or '').strip()
        name = (row.get('name', '') or '').strip()

        digits = re.sub(r'\D', '', phone)
        if len(digits) != 12 or not digits.startswith('998'):
            continue

        password = _gen_password()
        existing = ClientUser.objects.filter(phone=phone).first()

        if existing:
            if not include_existing:
                continue
            existing.set_password(password)
            existing.save(update_fields=['password_hash'])
            users_out.append({
                'id': existing.id,
                'name': existing.full_name,
                'phone': existing.phone,
                'org': existing.organization or '',
                'username': existing.username,
                'password': password,
                'status': _get_card_status(existing),
                'is_new': False,
            })
            updated += 1
        else:
            username = ClientUser.generate_username(name)
            user = ClientUser(
                phone=phone,
                username=username,
                full_name=name,
                organization=org,
            )
            user.set_password(password)
            user.is_verified = True
            if target_level:
                user.vip_level = target_level
            user.save()

            # Katta ERP Client bilan bog'lash
            try:
                from clients.models import Client
                from django.db.models import Q
                digits = re.sub(r'\D', '', phone)[-9:]
                client = Client.objects.filter(
                    Q(phone__endswith=digits) | Q(phone2__endswith=digits)
                ).first()
                if client:
                    user.client = client
                    user.save(update_fields=['client'])
            except Exception:
                pass

            users_out.append({
                'id': user.id,
                'name': user.full_name,
                'phone': user.phone,
                'org': user.organization or '',
                'username': user.username,
                'password': password,
                'status': _get_card_status(user),
                'is_new': True,
            })
            created += 1

    return JsonResponse({'ok': True, 'users': users_out, 'created': created, 'updated': updated})


@_staff_required
@require_POST
def mini_admin_card_users(request):
    """POST /mini/admin/card-users/ — tanlangan foydalanuvchilar ma'lumotlari."""
    body = json.loads(request.body)
    user_ids = body.get('user_ids', [])
    reset_password = body.get('reset_password', True)

    users = ClientUser.objects.filter(pk__in=user_ids).select_related('vip_level')
    users_out = []

    for user in users:
        password = None
        if reset_password:
            password = _gen_password()
            user.set_password(password)
            user.save(update_fields=['password_hash'])

        users_out.append({
            'id': user.id,
            'name': user.full_name,
            'phone': user.phone,
            'org': user.organization or '',
            'username': user.username,
            'password': password,
            'status': _get_card_status(user),
        })

    return JsonResponse({'ok': True, 'users': users_out})


# ═══════════════════════════════════════════════════════════════════
#  ANALITIKA
# ═══════════════════════════════════════════════════════════════════

@_staff_required
def mini_admin_analytics(request):
    """GET /mini/admin/analytics/ — login va faollik analitikasi."""
    from django.db.models import Count, Q
    from django.db.models.functions import TruncDate
    from django.utils import timezone
    from datetime import timedelta

    days = int(request.GET.get('days', 30))
    now = timezone.now()
    since = now - timedelta(days=days)

    attempts = ClientLoginAttempt.objects.filter(attempted_at__gte=since)

    # Umumiy statistikalar
    total = attempts.count()
    success = attempts.filter(success=True).count()
    failed = total - success

    by_type = attempts.filter(success=True).values('login_type').annotate(c=Count('id'))
    type_map = {r['login_type']: r['c'] for r in by_type}

    # Kunlik kirishlar (chart uchun)
    daily = (
        attempts.filter(success=True)
        .annotate(day=TruncDate('attempted_at'))
        .values('day')
        .annotate(
            total=Count('id'),
            web=Count('id', filter=Q(login_type='web')),
            telegram=Count('id', filter=Q(login_type='telegram')),
            qr=Count('id', filter=Q(login_type='qr')),
        )
        .order_by('day')
    )
    chart_labels = []
    chart_web = []
    chart_tg = []
    chart_qr = []
    for d in daily:
        chart_labels.append(d['day'].strftime('%d.%m'))
        chart_web.append(d['web'])
        chart_tg.append(d['telegram'])
        chart_qr.append(d['qr'])

    # Eng faol foydalanuvchilar
    top_users = (
        attempts.filter(success=True, user__isnull=False)
        .values('user__id', 'user__full_name', 'user__phone', 'user__username')
        .annotate(
            login_count=Count('id'),
            web_count=Count('id', filter=Q(login_type='web')),
            tg_count=Count('id', filter=Q(login_type='telegram')),
            qr_count=Count('id', filter=Q(login_type='qr')),
        )
        .order_by('-login_count')[:15]
    )

    # Oxirgi kirishlar (log)
    recent = attempts.filter(success=True).select_related('user').order_by('-attempted_at')[:50]
    recent_list = []
    for r in recent:
        recent_list.append({
            'phone': r.phone,
            'name': r.user.full_name if r.user else '—',
            'type': r.login_type,
            'ip': r.ip_address,
            'time': r.attempted_at.strftime('%d.%m.%Y %H:%M'),
        })

    # Muvaffaqiyatsiz urinishlar
    failed_list = []
    for r in attempts.filter(success=False).order_by('-attempted_at')[:30]:
        failed_list.append({
            'phone': r.phone,
            'type': r.login_type,
            'ip': r.ip_address,
            'time': r.attempted_at.strftime('%d.%m.%Y %H:%M'),
        })

    # Faol userlar (oxirgi 7 kun)
    week_ago = now - timedelta(days=7)
    active_week = ClientUser.objects.filter(last_activity__gte=week_ago).count()
    active_today = ClientUser.objects.filter(last_activity__date=now.date()).count()

    return JsonResponse({
        'ok': True,
        'summary': {
            'total_logins': success,
            'failed_attempts': failed,
            'web_logins': type_map.get('web', 0),
            'telegram_logins': type_map.get('telegram', 0),
            'qr_logins': type_map.get('qr', 0),
            'active_today': active_today,
            'active_week': active_week,
        },
        'chart': {
            'labels': chart_labels,
            'web': chart_web,
            'telegram': chart_tg,
            'qr': chart_qr,
        },
        'top_users': list(top_users),
        'recent': recent_list,
        'failed': failed_list,
    })


# ═══════════════════════════════════════════════════════════════════
#  TARIFLAR (ClientPlan) CRUD + foydalanuvchiga tayinlash
# ═══════════════════════════════════════════════════════════════════

def _safe_int(val, default=0):
    """Bo'sh/noto'g'ri qiymatni default'ga aylantiradi (500 oldini olish)."""
    try:
        if val in (None, ''):
            return default
        return int(val)
    except (ValueError, TypeError):
        return default


@_staff_required
def mini_admin_plans_data(request):
    """GET /mini/admin/plans/ — barcha tariflar + FEATURES ro'yxati (JSON)."""
    from ..services.features import FEATURES
    from ..services import limits as L

    plans = []
    for p in ClientPlan.objects.all().order_by('sort', 'id'):
        plans.append({
            'id': p.id,
            'name': p.name,
            'slug': p.slug,
            'price_uzs': p.price_uzs,
            'period_days': p.period_days,
            'coin_grant': p.coin_grant,
            'feature_keys': p.feature_keys or [],
            'limits': p.limits if isinstance(p.limits, dict) else {},
            'ai_included': p.ai_included,
            'is_free': p.is_free,
            'is_active': p.is_active,
            'sort': p.sort,
            'color': p.color,
            'icon': p.icon,
            'description': p.description,
            'user_count': p.users.count(),
        })

    features = [
        {'key': k, 'label': v['label'], 'category': v['category']}
        for k, v in FEATURES.items()
    ]
    limit_defs = [
        {'key': k, 'label': v['label'], 'unit': v['unit']}
        for k, v in L.LIMITS.items()
    ]
    return JsonResponse({'ok': True, 'plans': plans, 'features': features,
                         'limit_defs': limit_defs})


@_staff_required
@require_POST
def mini_admin_save_plan(request):
    """POST /mini/admin/save-plan/ — tarif yaratish/yangilash.

    is_free=True bo'lsa qolgan barcha tariflarning is_free=False qilinadi
    (faqat bitta standart/bepul tarif bo'lishi kerak).
    """
    from django.utils.text import slugify

    body = json.loads(request.body)
    plan_id = body.get('id')
    name = (body.get('name') or '').strip()
    if not name:
        return JsonResponse({'ok': False, 'error': 'Nomi majburiy'}, status=400)

    if plan_id:
        plan = get_object_or_404(ClientPlan, pk=plan_id)
    else:
        plan = ClientPlan()

    # Slug — berilsa ishlatamiz, bo'lmasa nomdan; noyoblik ta'minlanadi
    slug = (body.get('slug') or '').strip()
    if not slug:
        slug = slugify(name, allow_unicode=False) or 'plan'
    exclude_pk = plan.pk

    def _slug_taken(s):
        q = ClientPlan.objects.filter(slug=s)
        if exclude_pk:
            q = q.exclude(pk=exclude_pk)
        return q.exists()

    base_slug, i = slug, 2
    while _slug_taken(slug):
        slug = f'{base_slug}-{i}'
        i += 1
    plan.slug = slug

    plan.name = name
    plan.price_uzs = _safe_int(body.get('price_uzs'), 0)
    plan.period_days = _safe_int(body.get('period_days'), 30)
    plan.coin_grant = _safe_int(body.get('coin_grant'), 0)

    fk = body.get('feature_keys') or []
    plan.feature_keys = fk if isinstance(fk, list) else []

    # Miqdoriy limitlar (kvota). {key:int}. Faqat butun >=0 saqlanadi; blank/manfiy
    # = cheksiz (kalit qo'shilmaydi). Noma'lum kalitlar e'tiborsiz (registrga mos).
    from ..services import limits as L
    lim_in = body.get('limits') or {}
    limits_clean = {}
    if isinstance(lim_in, dict):
        for k, v in lim_in.items():
            if k not in L.LIMITS:
                continue
            try:
                iv = int(v)
            except (TypeError, ValueError):
                continue
            if iv >= 0:
                limits_clean[k] = iv
    plan.limits = limits_clean

    plan.ai_included = bool(body.get('ai_included', False))
    plan.is_free = bool(body.get('is_free', False))
    plan.is_active = bool(body.get('is_active', True))
    plan.sort = _safe_int(body.get('sort'), 0)
    plan.color = (body.get('color') or '').strip()
    plan.icon = (body.get('icon') or '').strip()
    plan.description = (body.get('description') or '').strip()
    plan.save()

    # Faqat bitta standart (bepul) tarif bo'lsin
    if plan.is_free:
        ClientPlan.objects.exclude(pk=plan.pk).filter(is_free=True).update(is_free=False)

    return JsonResponse({'ok': True, 'id': plan.id})


@_staff_required
@require_POST
def mini_admin_delete_plan(request, pk):
    """POST /mini/admin/delete-plan/<pk>/ — foydalanuvchisi/obunasi bor bo'lsa o'chirilmaydi."""
    from django.db.models import ProtectedError

    plan = get_object_or_404(ClientPlan, pk=pk)
    if plan.users.exists():
        return JsonResponse(
            {'ok': False, 'error': "Bu tarifda foydalanuvchi(lar) bor — avval ularni boshqa tarifga o'tkazing"},
            status=400,
        )
    try:
        plan.delete()
    except ProtectedError:
        return JsonResponse(
            {'ok': False, 'error': "Bu tarifga bog'liq obunalar bor — o'chirib bo'lmaydi"},
            status=400,
        )
    return JsonResponse({'ok': True})


@_staff_required
@require_POST
def mini_admin_assign_plan(request):
    """POST /mini/admin/assign-plan/ — foydalanuvchiga tarif tayinlash yoki olib tashlash.

    Body: {user_id, plan_id (bo'sh=olib tashlash), days (bo'sh=tarif davri)}.
    Tayinlanganda ClientSubscription (source='admin') tarix yozuvi yaratiladi.
    """
    from django.utils import timezone
    from datetime import timedelta

    body = json.loads(request.body)
    user_id = body.get('user_id')
    plan_id = body.get('plan_id')  # bo'sh/None = tarifni olib tashlash
    if not user_id:
        return JsonResponse({'ok': False, 'error': 'Foydalanuvchi tanlanmagan'}, status=400)

    user = get_object_or_404(ClientUser, pk=user_id)

    # Tarifni olib tashlash
    if not plan_id:
        user.plan = None
        user.plan_expires_at = None
        user.plan_since = None
        user.save(update_fields=['plan', 'plan_expires_at', 'plan_since'])
        return JsonResponse({'ok': True, 'removed': True})

    plan = get_object_or_404(ClientPlan, pk=plan_id)
    now = timezone.now()
    days = _safe_int(body.get('days'), 0)
    if days <= 0:
        days = plan.period_days or 30
    expires = now + timedelta(days=days)

    user.plan = plan
    user.plan_since = now
    user.plan_expires_at = expires
    user.save(update_fields=['plan', 'plan_since', 'plan_expires_at'])

    # Obuna tarix yozuvi
    ClientSubscription.objects.create(
        user=user,
        plan=plan,
        status=ClientSubscription.Status.ACTIVE,
        period_start=now,
        period_end=expires,
        source='admin',
    )
    return JsonResponse({'ok': True, 'expires': expires.strftime('%d.%m.%Y')})


# ═══════════════════════════════════════════════════════════════════
#  🪙 TANGA (CoinPack + AI narx + qo'lda berish)
# ═══════════════════════════════════════════════════════════════════

# Admin panelda ko'rsatiladigan standart AI amallar (coins.DEFAULT_PRICES bilan
# mos). ClientAIPrice yozuvi bo'lmasa shu default qiymatlar ko'rsatiladi.
_COIN_ACTIONS = [
    ('analytics_ai', 'AI Analitika'),
    ('ai_image',     'AI Rasm tahrirlash'),
    ('ai_panorama',  'AI Panorama (VR)'),
    ('laylo_chat',   'Laylo AI Chat'),
    ('gemini_live',  'Gemini Live'),
]


@_staff_required
def mini_admin_coins_data(request):
    """GET /mini/admin/coins/ — tanga paketlari + AI amal narxlari (JSON)."""
    from ..services import coins

    packs = []
    for p in CoinPack.objects.all().order_by('sort', 'id'):
        packs.append({
            'id': p.id,
            'name': p.name,
            'coins': p.coins,
            'bonus_coins': p.bonus_coins,
            'total_coins': p.total_coins,
            'price_uzs': p.price_uzs,
            'is_active': p.is_active,
            'sort': p.sort,
            'color': p.color,
        })

    # AI narxlari — 5 standart amal (DB yozuvi bo'lsa undan, bo'lmasa default)
    rows = {r.action_key: r for r in ClientAIPrice.objects.all()}
    prices = []
    for key, label in _COIN_ACTIONS:
        row = rows.get(key)
        if row is not None:
            prices.append({
                'action_key': key,
                'label': row.label or label,
                'coin_cost': row.coin_cost,
                'is_active': row.is_active,
                'exists': True,
            })
        else:
            prices.append({
                'action_key': key,
                'label': label,
                'coin_cost': int(coins.DEFAULT_PRICES.get(key, 0)),
                'is_active': True,
                'exists': False,
            })

    return JsonResponse({
        'ok': True,
        'packs': packs,
        'prices': prices,
        'coins_enabled': coins.ai_coins_enabled(),
    })


@_staff_required
@require_POST
def mini_admin_save_pack(request):
    """POST /mini/admin/save-pack/ — tanga paketi yaratish/yangilash."""
    body = json.loads(request.body)
    pack_id = body.get('id')
    name = (body.get('name') or '').strip()
    if not name:
        return JsonResponse({'ok': False, 'error': 'Nomi majburiy'}, status=400)

    if pack_id:
        pack = get_object_or_404(CoinPack, pk=pack_id)
    else:
        pack = CoinPack()

    pack.name = name
    pack.coins = _safe_int(body.get('coins'), 0)
    pack.bonus_coins = _safe_int(body.get('bonus_coins'), 0)
    pack.price_uzs = _safe_int(body.get('price_uzs'), 0)
    pack.is_active = bool(body.get('is_active', True))
    pack.sort = _safe_int(body.get('sort'), 0)
    pack.color = (body.get('color') or '').strip()
    pack.save()
    return JsonResponse({'ok': True, 'id': pack.id})


@_staff_required
@require_POST
def mini_admin_delete_pack(request, pk):
    """POST /mini/admin/delete-pack/<pk>/ — tanga paketini o'chirish."""
    pack = get_object_or_404(CoinPack, pk=pk)
    pack.delete()
    return JsonResponse({'ok': True})


@_staff_required
@require_POST
def mini_admin_save_aiprice(request):
    """POST /mini/admin/save-aiprice/ — AI amal narxini o'rnatish/yangilash."""
    body = json.loads(request.body)
    action_key = (body.get('action_key') or '').strip()
    if not action_key:
        return JsonResponse({'ok': False, 'error': 'action_key majburiy'}, status=400)

    ClientAIPrice.objects.update_or_create(
        action_key=action_key,
        defaults={
            'coin_cost': _safe_int(body.get('coin_cost'), 0),
            'label': (body.get('label') or '').strip(),
            'is_active': bool(body.get('is_active', True)),
        },
    )
    return JsonResponse({'ok': True})


@_staff_required
@require_POST
def mini_admin_grant_coins(request):
    """POST /mini/admin/grant-coins/ — foydalanuvchiga qo'lda tanga berish."""
    from ..services import coins

    body = json.loads(request.body)
    user_id = body.get('user_id')
    amount = _safe_int(body.get('amount'), 0)
    reason = (body.get('reason') or '').strip()
    if not user_id:
        return JsonResponse({'ok': False, 'error': 'Foydalanuvchi tanlanmagan'}, status=400)
    if amount <= 0:
        return JsonResponse({'ok': False, 'error': "Miqdor musbat bo'lishi kerak"}, status=400)

    user = get_object_or_404(ClientUser, pk=user_id)
    granted = coins.grant(user, amount, reason=reason or 'admin grant')
    return JsonResponse({'ok': True, 'granted': granted, 'balance': user.coins})


# ═══════════════════════════════════════════════════════════════════
#  💸 TO'LOVLAR (ClientPayment — faqat o'qish + sandbox qo'lda tasdiq)
# ═══════════════════════════════════════════════════════════════════

@_staff_required
def mini_admin_payments_data(request):
    """GET /mini/admin/payments/ — oxirgi 100 to'lov (JSON, faqat o'qish).

    `billing_live` bayrog'i qaytariladi: False (sandbox) bo'lsa UI'da qo'lda
    tasdiqlash tugmasi ko'rinadi, True (real) bo'lsa yo'q.
    """
    from ..models import ClientPayment, PaymeMerchantTxn
    from client_erp.payments import billing_live

    # Payme Merchant tranzaksiyalari (kassa callback holati) — order_id bo'yicha
    _TXN_STATE = {
        1: 'Yaratilgan (kutilmoqda)', 2: "To'langan",
        -1: 'Bekor (yaratilgandan keyin)', -2: "Bekor (to'langandan keyin)",
    }
    _txn_map = {}
    for t in PaymeMerchantTxn.objects.all().order_by('payment_id', '-pk'):
        if t.payment_id not in _txn_map:   # har order uchun eng oxirgisi
            _txn_map[t.payment_id] = {
                'payme_id': t.payme_id, 'state': t.state,
                'state_label': _TXN_STATE.get(t.state, str(t.state)),
                'amount_tiyin': t.amount_tiyin, 'reason': t.reason,
            }

    payments = []
    for p in ClientPayment.objects.select_related('user').all()[:100]:
        u = p.user
        payments.append({
            'id': p.id,
            'order_id': p.id,                       # Payme account.order_id — AYNAN shu
            'amount_tiyin': int(p.amount_uzs or 0) * 100,
            'checkout_url': p.checkout_url or '',
            'payme_txn': _txn_map.get(p.id),        # None = Payme hali murojaat qilmagan
            'user_id': u.id if u else None,
            'user_name': (u.full_name if u else '') or (u.username if u else '') or '—',
            'username': u.username if u else '',
            'provider': p.provider,
            'provider_label': p.get_provider_display(),
            'purpose': p.purpose,
            'purpose_label': p.get_purpose_display(),
            'target_id': p.target_id,
            'amount_uzs': p.amount_uzs,
            'status': p.status,
            'status_label': p.get_status_display(),
            'fulfilled': p.fulfilled,
            'created_at': p.created_at.strftime('%d.%m.%Y %H:%M') if p.created_at else '',
            'paid_at': p.paid_at.strftime('%d.%m.%Y %H:%M') if p.paid_at else '',
        })

    return JsonResponse({'ok': True, 'payments': payments, 'billing_live': billing_live()})


@_staff_required
@require_POST
def mini_admin_confirm_payment(request):
    """POST /mini/admin/confirm-payment/ — sandbox to'lovni QO'LDA tasdiqlash.

    XAVFSIZLIK INVARIANTI: FAQAT sandbox rejimida (billing_live()=False)
    ishlaydi. Real rejimda qo'lda tasdiq BLOKLANADI — pulni faqat provayder
    webhook'i tasdiqlashi kerak. service.confirm() idempotent (fulfilled bayrog'i
    orqali) — takroriy tasdiq qayta tanga bermaydi / tarifni ikki marta bermaydi.
    """
    from ..models import ClientPayment
    from client_erp.payments import service, billing_live

    if billing_live():
        return JsonResponse({'ok': False, 'error': "Real rejimda qo'lda tasdiq yo'q"}, status=400)

    body = json.loads(request.body or '{}')
    payment_id = body.get('payment_id')
    if not payment_id:
        return JsonResponse({'ok': False, 'error': 'payment_id majburiy'}, status=400)

    p = get_object_or_404(ClientPayment, pk=payment_id)
    service.confirm(p)
    return JsonResponse({'ok': True, 'status': p.status, 'fulfilled': p.fulfilled})


@_staff_required
def mini_admin_export_users(request):
    """GET /mini/admin/export-users/?period=week|3day|all — userlar ro'yxatini
    Excel (.xlsx) qilib qaytaradi (2026-07-16).

    Faqat O'QISH: ClientUser'dan tanlab, xlsx yasaydi. Moliyaga aloqasi yo'q.
    Himoya: mini_admin_dashboard bilan bir xil (@_staff_required).
    """
    from django.http import HttpResponse
    from django.utils import timezone as _tz
    from datetime import timedelta as _td
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    period = (request.GET.get('period') or 'week').strip()
    qs = ClientUser.objects.order_by('-created_at')
    today = _tz.localdate()
    if period == '3day':
        qs = qs.filter(created_at__date__gte=today - _td(days=2)); label = 'oxirgi-3-kun'
    elif period == 'all':
        label = 'hammasi'
    else:  # week (default)
        qs = qs.filter(created_at__date__gte=today - _td(days=6)); label = 'oxirgi-hafta'

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Userlar"
    headers = ["№", "Ism", "Telegram", "Telefon", "Username", "Ro'yxatdan o'tgan"]
    ws.append(headers)
    head_fill = PatternFill("solid", fgColor="16A34A")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
    for i, u in enumerate(qs, 1):
        ws.append([
            i,
            u.full_name or '',
            ('@' + u.telegram_username) if u.telegram_username else '',
            u.phone or '',
            u.username or '',
            _tz.localtime(u.created_at).strftime('%Y-%m-%d %H:%M') if u.created_at else '',
        ])
    # Ustun kengligi
    for col, w in zip('ABCDEF', (5, 26, 22, 16, 18, 18)):
        ws.column_dimensions[col].width = w

    resp = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    fname = f"bittada-usta-userlar-{label}-{today.strftime('%Y%m%d')}.xlsx"
    resp['Content-Disposition'] = f'attachment; filename="{fname}"'
    wb.save(resp)
    return resp
